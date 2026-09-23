from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "documentation" / "GUIDE_UTILISATEUR_CLEANGO.md"
OUTPUT = ROOT / "documentation" / "GUIDE_UTILISATEUR_CLEANGO.docx"
BRAND = RGBColor(176, 74, 24)
DARK = RGBColor(31, 41, 55)
MUTED = RGBColor(92, 102, 115)


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Dans Word : clic droit ici, puis « Mettre à jour les champs »."
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = BRAND
        else:
            paragraph.add_run(part)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = DARK
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Title", 28), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11)):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = BRAND
    styles[name].font.bold = True

# Couverture
doc.add_paragraph().add_run("CLEANGO").font.color.rgb = BRAND
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Guide utilisateur\net support de formation")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Gestion opérationnelle d’une station de lavage automobile")
run.font.size = Pt(14)
run.font.color.rgb = MUTED
doc.add_paragraph()
cover_image = ROOT / "documentation" / "captures" / "02-tableau-de-bord.png"
doc.add_picture(str(cover_image), width=Inches(6.4))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("\nVersion 1.0  •  25 juillet 2026\n").bold = True
meta.add_run("Responsables, caissiers, agents d’accueil, superviseurs et administrateurs")
doc.add_page_break()

doc.add_heading("Sommaire", level=1)
add_toc(doc.add_paragraph())
doc.add_page_break()

lines = SOURCE.read_text(encoding="utf-8").splitlines()
i = 0
skip_intro_metadata = True
while i < len(lines):
    line = lines[i].strip()
    if line == "---":
        skip_intro_metadata = False
        i += 1
        continue
    if skip_intro_metadata:
        i += 1
        continue
    if not line:
        i += 1
        continue

    image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
    if image:
        caption, relative = image.groups()
        image_path = SOURCE.parent / relative
        doc.add_picture(str(image_path), width=Inches(6.45))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        i += 1
        continue

    heading = re.match(r"^(#{1,3})\s+(.+)", line)
    if heading:
        level = len(heading.group(1))
        title_text = heading.group(2)
        if title_text != "CleanGo — Guide utilisateur et support de formation":
            if level == 2 and title_text.startswith(("6.", "7.", "8.", "9.", "10.", "11.", "12.", "13.", "14.", "15.", "16.", "17.", "18.", "19.", "20.")):
                doc.add_page_break()
            doc.add_heading(title_text, level=level)
        i += 1
        continue

    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        headers = [c.strip() for c in line.strip("|").split("|")]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = True
        for idx, value in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = value
            shade_cell(cell, "B04A18")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
        set_repeat_table_header(table.rows[0])
        for row_values in rows:
            cells = table.add_row().cells
            for idx, value in enumerate(row_values):
                cells[idx].text = value
                cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        doc.add_paragraph()
        continue

    bullet = re.match(r"^-\s+(.+)", line)
    numbered = re.match(r"^\d+\.\s+(.+)", line)
    if bullet or numbered:
        p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
        add_inline(p, (bullet or numbered).group(1))
        i += 1
        continue

    p = doc.add_paragraph()
    add_inline(p, line)
    i += 1

# En-tête et pied de page
for section in doc.sections:
    header = section.header.paragraphs[0]
    header.text = "CleanGo — Guide utilisateur"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CleanGo  •  Documentation et formation  •  ")
    field_run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instr, end])
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

settings = doc.settings._element
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)

doc.save(OUTPUT)
print(f"Document Word créé : {OUTPUT}")
